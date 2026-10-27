"""Generate the resume from the same verified profile data used by the portfolio.

Requires python-docx and a local Chrome or Edge installation. Run from any folder:
    python scripts/generate_resume.py
Pass --browser PATH to choose another Chromium-based browser.
Pass --pdf-source PATH to copy a supplied PDF unchanged instead of rendering one.
"""

import argparse
import html
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROFILE = json.loads((ROOT / "src/data/profile.json").read_text(encoding="utf-8"))
OUTPUT = ROOT / "public"
STEM = "Ronnel_Daryl_Rontos_Resume"


def escape(value):
    return html.escape(value, quote=True)


def role_html(role):
    bullets = "".join(f"<li>{escape(bullet)}</li>" for bullet in role["bullets"])
    return (
        '<article class="role">'
        f'<div class="role-heading"><h3>{escape(role["role"])}</h3>'
        f'<span>{escape(role["period"])}</span></div>'
        f'<p class="company">{escape(role["company"])}</p>'
        f"<ul>{bullets}</ul></article>"
    )


def generate_html():
    skills = "".join(
        f'<p><strong>{escape(category)}:</strong> {escape(", ".join(items))}</p>'
        for category, items in PROFILE["skills"].items()
    )
    roles = [role_html(role) for role in PROFILE["experiences"]]
    projects = "".join(
        f'<p><strong>{escape(project["name"])}:</strong> {escape(project["description"])}</p>'
        for project in PROFILE["projects"]
    )
    education = PROFILE["education"]
    content = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(PROFILE['name'])} — Resume</title>
<style>
@page {{ size: A4; margin: 14mm 16mm; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #202b2b; font: 10pt/1.38 Arial, Helvetica, sans-serif; }}
a {{ color: #245b4a; text-decoration: none; }}
header {{ border-top: 3pt solid #245b4a; padding-top: 12pt; margin-bottom: 12pt; }}
h1 {{ margin: 0; font-size: 24pt; line-height: 1.15; letter-spacing: .6pt; color: #163d33; }}
.title {{ font-size: 11.2pt; margin: 5pt 0 7pt; color: #245b4a; }}
.contact, .links {{ font-size: 9pt; margin: 3pt 0; }}
.links a + a {{ margin-left: 14pt; }}
h2 {{ font-size: 10pt; letter-spacing: 1.2pt; text-transform: uppercase; color: #245b4a; border-bottom: .6pt solid #bed0c8; padding-bottom: 4pt; margin: 13pt 0 7pt; break-after: avoid; }}
p {{ margin: 0 0 5pt; }}
.skills {{ font-size: 9.5pt; line-height: 1.38; }}
.skills p {{ margin-bottom: 4pt; }}
.role {{ margin-bottom: 10pt; break-inside: avoid; }}
.role-heading {{ display: flex; justify-content: space-between; gap: 12pt; align-items: baseline; }}
.role-heading h3 {{ margin: 0; font-size: 10.5pt; font-weight: bold; }}
.role-heading > span {{ flex-shrink: 0; font-size: 9pt; color: #465550; }}
.company {{ margin: 2pt 0 5pt; color: #245b4a; font-weight: bold; font-size: 10pt; }}
ul {{ margin: 0; padding-left: 13pt; }}
li {{ padding-left: 1pt; margin-bottom: 3pt; }}
.sheet + .sheet {{ break-before: page; }}
.continuation {{ display: flex; justify-content: space-between; align-items: baseline; border-top: 2pt solid #245b4a; padding-top: 9pt; margin-bottom: 13pt; font-size: 9pt; color: #465550; }}
.continuation strong {{ color: #163d33; font-size: 11pt; }}
.projects p {{ margin-bottom: 7pt; }}
.education {{ break-inside: avoid; }}
@media screen {{ body {{ background: #e9edeb; padding: 24px 12px; }} .sheet {{ background: white; max-width: 210mm; min-height: 297mm; margin: 0 auto 24px; padding: 14mm 16mm; box-shadow: 0 3px 20px #0001; }} }}
</style></head><body>
<main>
<section class="sheet" aria-label="Resume page 1">
<header>
<h1>{escape(PROFILE['name'].upper())}</h1>
<p class="title">{escape(PROFILE['title'])}</p>
<p class="contact">{escape(PROFILE['location'])} | {escape(PROFILE['phone'])} | <a href="mailto:{escape(PROFILE['email'])}">{escape(PROFILE['email'])}</a></p>
<p class="links"><a href="{escape(PROFILE['linkedin'])}">linkedin.com/in/darylrontos</a><a href="{escape(PROFILE['github'])}">github.com/GinTrue</a></p>
</header>
<h2>Professional summary</h2><p>{escape(PROFILE['summary'])}</p>
<h2>Technical skills</h2><div class="skills">{skills}</div>
<h2>Professional experience</h2>{roles[0]}{roles[1]}
</section>
<section class="sheet" aria-label="Resume page 2">
<div class="continuation"><strong>{escape(PROFILE['name'])}</strong><span>Professional experience, projects & education</span></div>
<h2>Professional experience — continued</h2>{roles[2]}{roles[3]}
<h2>Selected projects</h2><div class="projects">{projects}</div>
<div class="education"><h2>Education</h2><p><strong>{escape(education['degree'])}</strong><br>{escape(education['school'])} | {escape(education['period'])}</p></div>
</section>
</main></body></html>
"""
    path = OUTPUT / f"{STEM}.html"
    path.write_text(content, encoding="utf-8")
    return path


def generate_docx():
    document = Document()
    document.core_properties.title = f"{PROFILE['name']} — Resume"
    document.core_properties.author = PROFILE["name"]
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.63)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("202B2B")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for name in ("Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string("245B4A")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    bullet_style = document.styles["List Bullet"]
    bullet_style.paragraph_format.space_after = Pt(3)
    bullet_style.paragraph_format.line_spacing = 1.12

    def paragraph(text="", bold=False):
        item = document.add_paragraph()
        item.add_run(text).bold = bold
        return item

    def heading(text):
        document.add_heading(text.upper(), level=1)

    def role(item):
        title = paragraph(item["role"], bold=True)
        title.paragraph_format.keep_with_next = True
        company = paragraph(f'{item["company"]} | {item["period"]}')
        company.paragraph_format.keep_with_next = True
        for index, bullet in enumerate(item["bullets"]):
            line = document.add_paragraph(bullet, style="List Bullet")
            line.paragraph_format.keep_with_next = index < len(item["bullets"]) - 1

    name = paragraph(PROFILE["name"].upper(), bold=True)
    name.runs[0].font.size = Pt(23)
    name.runs[0].font.color.rgb = RGBColor.from_string("163D33")
    paragraph(PROFILE["title"])
    contact = paragraph(f'{PROFILE["location"]} | {PROFILE["phone"]} | {PROFILE["email"]}')
    contact.runs[0].font.size = Pt(9)
    links = paragraph("linkedin.com/in/darylrontos | github.com/GinTrue")
    links.runs[0].font.size = Pt(9)
    heading("Professional summary")
    paragraph(PROFILE["summary"])
    heading("Technical skills")
    for category, items in PROFILE["skills"].items():
        line = document.add_paragraph()
        line.add_run(f"{category}: ").bold = True
        line.add_run(", ".join(items))
        for run in line.runs:
            run.font.size = Pt(9.5)
    heading("Professional experience")
    for item in PROFILE["experiences"][:2]:
        role(item)
    document.add_page_break()
    paragraph(PROFILE["name"], bold=True)
    heading("Professional experience — continued")
    for item in PROFILE["experiences"][2:]:
        role(item)
    heading("Selected projects")
    for project in PROFILE["projects"]:
        line = document.add_paragraph()
        line.add_run(f'{project["name"]}: ').bold = True
        line.add_run(project["description"])
    heading("Education")
    paragraph(PROFILE["education"]["degree"], bold=True)
    paragraph(f'{PROFILE["education"]["school"]} | {PROFILE["education"]["period"]}')
    # Keep the footer separate from the resume's reading order.
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    page_number = OxmlElement("w:fldSimple")
    page_number.set(qn("w:instr"), "PAGE")
    footer._p.append(page_number)
    path = OUTPUT / f"{STEM}.docx"
    document.save(path)
    return path


def find_browser(explicit):
    if explicit:
        return str(Path(explicit).resolve(strict=True))
    candidates = [shutil.which(name) for name in ("chromium", "google-chrome", "chrome", "msedge")]
    for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        base = os.environ.get(variable)
        if base:
            candidates += [str(Path(base) / suffix) for suffix in (
                "Google/Chrome/Application/chrome.exe", "Microsoft/Edge/Application/msedge.exe"
            )]
    return next((path for path in candidates if path and Path(path).is_file()), None)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--browser", help="Path to a Chromium-based browser")
    parser.add_argument("--pdf-source", type=Path, help="Use an existing PDF unchanged")
    args = parser.parse_args()
    source_pdf = args.pdf_source.resolve(strict=True) if args.pdf_source else None
    if source_pdf:
        with source_pdf.open("rb") as supplied_pdf:
            if supplied_pdf.read(5) != b"%PDF-":
                parser.error("--pdf-source must point to a PDF file.")
    browser = find_browser(args.browser) if not source_pdf else None
    if not source_pdf and not browser:
        parser.error("Chrome or Edge is required for PDF export; specify --browser PATH.")
    OUTPUT.mkdir(exist_ok=True)
    source = generate_html()
    word = generate_docx()
    pdf = OUTPUT / f"{STEM}.pdf"
    if source_pdf:
        if source_pdf != pdf.resolve():
            shutil.copyfile(source_pdf, pdf)
    else:
        with tempfile.TemporaryDirectory(prefix="rontos-resume-") as temporary:
            temporary_pdf = Path(temporary) / f"{STEM}.pdf"
            result = subprocess.run([
                browser, "--headless", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
                "--disable-extensions", f"--user-data-dir={Path(temporary) / 'browser'}",
                "--no-pdf-header-footer", f"--print-to-pdf={temporary_pdf}", source.as_uri(),
            ], capture_output=True, text=True, timeout=60)
            if result.returncode or not temporary_pdf.is_file():
                raise RuntimeError(f"Resume PDF export failed: {result.stderr}")
            shutil.copyfile(temporary_pdf, pdf)
    for path in (pdf, word, source):
        print(f"Created {path.relative_to(ROOT)} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
